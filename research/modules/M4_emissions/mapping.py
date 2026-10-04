"""M4 helpers: NAICS 2017 -> SIC 1987 -> Fama-French 49 mapping and supply-chain GHG intensity by FF49 industry.

Sources (all in data/raw/):
  epa_sc_ghg_naics_v13.csv            EPA Supply Chain GHG Emission Factors v1.3 (kg CO2e per 2022 USD, purchaser price,
                                      2017 NAICS, 6-digit; electricity 2211xx is NOT in the file by EPA design)
  census_1987_SIC_to_2002_NAICS.xls   Census concordance, 1987 SIC -> 2002 NAICS
  census_2002_to_2007_NAICS.xls       Census concordance, 2002 -> 2007 NAICS
  census_2007_to_2012_NAICS.xls       Census concordance, 2007 -> 2012 NAICS
  census_2012_to_2017_NAICS.xlsx      Census concordance, 2012 -> 2017 NAICS
  useeio201_ghg_DN.csv                USEEIO v2.0.1-411 (EPA) GHG row of the D (direct) and N (total) matrices,
                                      kg CO2e per 2012 USD, extracted from USEEIOv2.0.1-411.xlsx (see extract_useeio()).
"""
from __future__ import annotations
import pathlib, sys
import numpy as np
import pandas as pd

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "lib"))
from common import RAW, DERIVED, kf_industry_sic_map  # noqa: E402

USEEIO_URL = "https://pasteur.epa.gov/uploads/10.23719/1524311/USEEIOv2.0.1-411.xlsx"
ELEC_NAICS17 = [221111, 221112, 221113, 221114, 221115, 221116, 221117, 221118, 221121, 221122]


# ----------------------------------------------------------------------------- raw loaders
def load_epa() -> pd.DataFrame:
    e = pd.read_csv(RAW / "epa_sc_ghg_naics_v13.csv")
    e.columns = ["naics", "title", "ghg", "unit", "sef_nomargin", "margin", "sef", "useeio"]
    e["naics"] = e["naics"].astype(int)
    e["title"] = e["title"].str.strip()
    return e[["naics", "title", "sef", "sef_nomargin", "margin", "useeio"]]


def _step(fname: str) -> pd.DataFrame:
    x = pd.read_excel(RAW / fname, header=2, dtype=str).iloc[:, :4]
    x.columns = ["old", "old_title", "new", "new_title"]
    x = x[x["old"].str.fullmatch(r"\d{6}", na=False) & x["new"].str.fullmatch(r"\d{6}", na=False)]
    return x.assign(old=x["old"].astype(int), new=x["new"].astype(int))[["old", "new", "new_title"]]


def load_sic_naics02() -> pd.DataFrame:
    d = pd.read_excel(RAW / "census_1987_SIC_to_2002_NAICS.xls", header=0, dtype=str)
    d.columns = ["sic", "sic_title", "naics02", "naics02_title"]
    d = d[d["sic"].str.fullmatch(r"\d{3,4}", na=False) & d["naics02"].str.fullmatch(r"\d{6}", na=False)]
    return d.assign(sic=d["sic"].astype(int), naics02=d["naics02"].astype(int))[["sic", "sic_title", "naics02"]]


def sic_to_ff49(sic: int, sicmap=None) -> str | None:
    sicmap = sicmap or kf_industry_sic_map()
    for ind, ranges in sicmap.items():
        for lo, hi in ranges:
            if lo <= sic <= hi:
                return ind
    return None


# ----------------------------------------------------------------------------- chain NAICS 2017 -> SIC 1987
def naics17_sic_links() -> pd.DataFrame:
    """Every (naics17, sic) pair reachable through the Census chain 2017<-2012<-2007<-2002<-SIC87."""
    s12 = _step("census_2012_to_2017_NAICS.xlsx").rename(columns={"old": "naics12", "new": "naics17", "new_title": "naics17_title"})
    s07 = _step("census_2007_to_2012_NAICS.xls").rename(columns={"old": "naics07", "new": "naics12"})[["naics07", "naics12"]]
    s02 = _step("census_2002_to_2007_NAICS.xls").rename(columns={"old": "naics02", "new": "naics07"})[["naics02", "naics07"]]
    sic = load_sic_naics02()
    ch = s12.merge(s07, on="naics12", how="left").merge(s02, on="naics07", how="left").merge(sic, on="naics02", how="left")
    ch = ch.dropna(subset=["sic"]).assign(sic=lambda d: d["sic"].astype(int))
    titles = s12.drop_duplicates("naics17").set_index("naics17")["naics17_title"]
    man = pd.DataFrame([{"naics17": k, "naics17_title": titles.get(k, ""), "sic": v[0], "sic_title": "MANUAL: " + v[1]}
                        for k, v in MANUAL_LINKS.items()])
    ch = pd.concat([ch, man], ignore_index=True)
    return ch[["naics17", "naics17_title", "naics12", "naics07", "naics02", "sic", "sic_title"]].drop_duplicates()


# EPA NAICS codes the Census chain leaves without a numeric SIC predecessor: manual links (documented judgment).
MANUAL_LINKS = {
    112130: (212, "Dual-purpose cattle ranching: Census lists no SIC; SIC 0212 beef cattle except feedlots (Agric)"),
    541120: (7389, "Offices of notaries: Census lists no SIC; SIC 7389 business services NEC (BusSv)"),
}
# 551114 (corporate, subsidiary and regional managing offices) maps only to SIC 'Aux' (auxiliary establishments): excluded.
# Documented overrides where the plurality rule misassigns a NAICS code: the 2017 R&D codes inherit SIC links from the
# R&D "pieces" of aircraft (3721/3724/3728) and guided-missile (3761/3764/3769) SIC codes, which tie Aero and Guns at
# 3 links each; commercial research SIC 8731/8733 (the core of these codes) is BusSv.
OVERRIDES = {
    541713: ("BusSv", "R&D in nanotechnology: plurality tie Aero/Guns from R&D pieces of 372x/376x; core SIC 8731/8733 is BusSv"),
    541714: ("BusSv", "R&D in biotechnology: same as 541713"),
    541715: ("BusSv", "R&D in physical, engineering and life sciences: same as 541713"),
}


def build_mapping() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return (links, naics_map).
    links: one row per (naics17, sic) with its FF49 industry.
    naics_map: one row per (naics17, ff49) under two rules:
      primary   = the FF49 industry holding the plurality of the NAICS code's distinct SIC predecessors (ties: all tied);
      any_link  = every FF49 industry with at least one SIC predecessor (many-to-many robustness)."""
    sicmap = kf_industry_sic_map()
    links = naics17_sic_links()
    uniq = links[["sic"]].drop_duplicates()
    uniq["ff49"] = [sic_to_ff49(s, sicmap) for s in uniq["sic"]]
    links = links.merge(uniq, on="sic", how="left")
    links["ff49"] = links["ff49"].fillna("UNASSIGNED")
    d = links[["naics17", "naics17_title", "sic", "ff49"]].drop_duplicates(["naics17", "sic"])
    cnt = d.groupby(["naics17", "ff49"]).size().rename("n_sic").reset_index()
    tot = d.groupby("naics17").size().rename("n_sic_total")
    cnt = cnt.merge(tot, on="naics17")
    cnt["share"] = cnt["n_sic"] / cnt["n_sic_total"]
    # plurality among assigned FF49 industries (UNASSIGNED SIC codes cannot win)
    assigned = cnt[cnt["ff49"] != "UNASSIGNED"].copy()
    mx = assigned.groupby("naics17")["n_sic"].transform("max")
    assigned["primary"] = assigned["n_sic"] == mx
    for code, (ind, why) in OVERRIDES.items():
        m = assigned["naics17"] == code
        if m.any():
            assigned.loc[m, "primary"] = assigned.loc[m, "ff49"] == ind
            if not (assigned.loc[m, "primary"]).any():   # override target had no SIC link: add it explicitly
                assigned = pd.concat([assigned, pd.DataFrame([{"naics17": code, "ff49": ind, "n_sic": 0, "n_sic_total": int(tot.get(code, 0)),
                                                               "share": 0.0, "primary": True}])], ignore_index=True)
    assigned["override_note"] = assigned["naics17"].map({k: v[1] for k, v in OVERRIDES.items()}).fillna("")
    titles = links.drop_duplicates("naics17").set_index("naics17")["naics17_title"]
    assigned["naics17_title"] = assigned["naics17"].map(titles)
    assigned["any_link"] = True
    return links, assigned.sort_values(["naics17", "n_sic"], ascending=[True, False]).reset_index(drop=True)


# ----------------------------------------------------------------------------- USEEIO direct / total and electricity patch
def extract_useeio(xlsx_path: pathlib.Path) -> pd.DataFrame:
    """Pull the 'Greenhouse Gases' row of D (direct, per $ output) and N (total, per $ commodity) from USEEIOv2.0.1-411.xlsx."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    out = {}
    for sh in ("D", "N"):
        rows = list(wb[sh].iter_rows(values_only=True))
        df = pd.DataFrame(rows[1:], columns=["indicator"] + list(rows[0][1:])).set_index("indicator")
        out[sh] = df.loc["Greenhouse Gases"].astype(float)
    rows = list(wb["B"].iter_rows(values_only=True))          # direct flows per $ output; CO2 row only
    df = pd.DataFrame(rows[1:], columns=["flow"] + list(rows[0][1:])).set_index("flow")
    out["B_CO2"] = df.loc["Carbon dioxide/emission/air/kg"].astype(float)
    res = pd.DataFrame(out)
    res.index = res.index.str.replace("/US", "", regex=False)
    res.index.name = "useeio"
    return res.rename(columns={"D": "useeio_direct", "N": "useeio_total", "B_CO2": "useeio_direct_co2"})


def load_useeio() -> pd.DataFrame:
    p = RAW / "useeio201_ghg_DN.csv"
    if not p.exists():
        import requests
        big = RAW / "USEEIOv2.0.1-411.xlsx"
        if not big.exists():
            big.write_bytes(requests.get(USEEIO_URL, timeout=600).content)
        extract_useeio(big).to_csv(p)
    return pd.read_csv(p, index_col=0)


def _useeio_codes(codes: str, u: pd.DataFrame) -> tuple[list, str]:
    """USEEIO v2.0.1 sector codes for a v1.3 reference code string. Exact match first; otherwise (v2.2 re-aggregated
    codes such as 562000 or 335220) all v2.0.1 codes sharing the longest common prefix of at least 3 characters."""
    out, how = [], "exact"
    for c in [c.strip() for c in str(codes).split(",")]:
        if c in u.index:
            out.append(c); continue
        for k in range(len(c) - 1, 2, -1):
            hit = [x for x in u.index if x.startswith(c[:k])]
            if hit:
                out += hit; how = f"prefix{k}"; break
    return out, how


def _useeio_lookup(codes: str, u: pd.DataFrame, col: str) -> float:
    cs, _ = _useeio_codes(codes, u)
    return float(np.mean([u.at[c, col] for c in cs])) if cs else np.nan


def naics_factors() -> tuple[pd.DataFrame, dict]:
    """EPA v1.3 factors per NAICS 2017 plus USEEIO v2.0.1 direct/total, with the electricity patch appended.
    Electricity patch: EPA excludes NAICS 2211xx. I use USEEIO v2.0.1 total (N) for 221100, rescaled to v1.3 units by the
    median ratio EPA_v1.3_without_margins / USEEIO_N across all NAICS codes whose reference USEEIO code matches.
    Electricity carries no trade/transport margin, so with-margins = without-margins for these rows."""
    e = load_epa()
    u = load_useeio()
    e["useeio_direct"] = [_useeio_lookup(c, u, "useeio_direct") for c in e["useeio"]]
    e["useeio_direct_co2"] = [_useeio_lookup(c, u, "useeio_direct_co2") for c in e["useeio"]]
    e["useeio_total"] = [_useeio_lookup(c, u, "useeio_total") for c in e["useeio"]]
    e["useeio_match"] = [_useeio_codes(c, u)[1] for c in e["useeio"]]
    ratio = (e["sef_nomargin"] / e["useeio_total"]).replace([np.inf, -np.inf], np.nan).dropna()
    scale = float(ratio.median())
    elec_total = float(u.at["221100", "useeio_total"]); elec_direct = float(u.at["221100", "useeio_direct"])
    patch_val = elec_total * scale
    titles = {221111: "Hydroelectric Power Generation", 221112: "Fossil Fuel Electric Power Generation",
              221113: "Nuclear Electric Power Generation", 221114: "Solar Electric Power Generation",
              221115: "Wind Electric Power Generation", 221116: "Geothermal Electric Power Generation",
              221117: "Biomass Electric Power Generation", 221118: "Other Electric Power Generation",
              221121: "Electric Bulk Power Transmission and Control", 221122: "Electric Power Distribution"}
    add = pd.DataFrame({"naics": ELEC_NAICS17, "title": [titles[c] for c in ELEC_NAICS17], "sef": patch_val,
                        "sef_nomargin": patch_val, "margin": 0.0, "useeio": "221100",
                        "useeio_direct": elec_direct, "useeio_direct_co2": float(u.at["221100", "useeio_direct_co2"]),
                        "useeio_total": elec_total, "useeio_match": "exact"})
    e["source"] = "EPA v1.3"; add["source"] = "patch: USEEIO v2.0.1 N(221100) x median(EPA v1.3 / USEEIO N)"
    info = {"scale_median": scale, "scale_iqr_lo": float(ratio.quantile(0.25)), "scale_iqr_hi": float(ratio.quantile(0.75)),
            "n_ratio": int(len(ratio)), "elec_useeio_total_2012usd": elec_total, "elec_useeio_direct_2012usd": elec_direct,
            "elec_patch_kg_per_2022usd": patch_val,
            "n_epa_codes_useeio_matched": int(e["useeio_total"].notna().sum()), "n_epa_codes": int(len(e))}
    return pd.concat([e, add], ignore_index=True), info


def aggregate_ff49(fac: pd.DataFrame, nmap: pd.DataFrame, rule: str = "primary") -> pd.DataFrame:
    """Unweighted mean/median of NAICS factors by FF49 industry (no output weights)."""
    m = nmap[nmap[rule]][["naics17", "ff49"]].merge(fac, left_on="naics17", right_on="naics", how="inner")
    g = m.groupby("ff49")
    return pd.DataFrame({"epa_sc_mean": g["sef"].mean(), "epa_sc_median": g["sef"].median(),
                         "epa_sc_nomargin_mean": g["sef_nomargin"].mean(), "epa_sc_nomargin_median": g["sef_nomargin"].median(),
                         "useeio_direct_mean": g["useeio_direct"].mean(), "useeio_direct_median": g["useeio_direct"].median(),
                         "useeio_direct_co2_mean": g["useeio_direct_co2"].mean(),
                         "useeio_total_mean": g["useeio_total"].mean(),
                         "n_naics": g["sef"].count()})
