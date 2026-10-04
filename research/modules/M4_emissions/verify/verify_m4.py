"""Independent verification of M4_emissions (does NOT import modules/M4_emissions/*.py).

Run: cd /home/hashim/projects/GA/project/research && uv run python modules/M4_emissions/verify/verify_m4.py

Rebuilds from raw files, with my own code:
  1. team-file fingerprints (Util/Fun ratio, tie groups, Aero/Ships ranks)
  2. NAICS-2017 -> SIC-1987 -> FF49 mapping, using the Census 2002-NAICS -> 1987-SIC file (the builder used the
     opposite-direction 1987-SIC -> 2002-NAICS file), my own Siccodes49 parser, plurality rule, electricity patch
     computed from the USEEIO v2.0.1 workbook directly (not from the builder's extracted CSV)
  3. EPA FF49 ranking, legs, Spearman vs team file, direct-GHG ranks
  4. Aero/Ships leave-two-out calibration t-stats
  5. Green-minus-Brown returns, FF5+UMD alpha with my own Newey-West (Bartlett, 6 lags, no df correction, normal p)
     for the primary test, the 9-variant family (post-2010 and full sample, Holm), raw performance, recent-window betas
  6. Energy-rally variants
  7. Team Short-Brown timing rule re-implemented from the team notebook, with the Brown leg swapped
  8. Ledger completeness and p-value consistency against the module's CSVs
  9. Extra sensitivity: the 5th Green member (Smoke vs Chips differ by 0.001 kg/$) and drop-one-industry
 10. Round 2 (after the builder's fixes F1-F11): calibration ratios + Holm/BH over 12 tests incl. direct CO2 read from
     USEEIO sheet B, rank/spread evidence, leg-membership table, CAPM beta decomposition, EPA-minus-team series,
     full-sample family wording, performance/energy/timing-rule numbers in the revised text, ledger completeness
     (every unadjusted p printed in an M4 table must be in the ledger). Round-2 checks are prefixed "R2".
Writes verify/verify_results.csv and verify/verify_sensitivity.csv (data only).
"""
from __future__ import annotations
import pathlib, re, sys
import numpy as np
import pandas as pd
from scipy import stats

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(ROOT / "lib"))
from common import RAW, TABLES, load_team, load_ff5_mom, load_kf_industries  # loaders only

pd.set_option("display.width", 220)
RES: list[dict] = []


def check(claim, reported, recomputed, tol, note=""):
    ok = (np.isfinite(reported) and np.isfinite(recomputed) and abs(reported - recomputed) <= tol)
    RES.append({"claim": claim, "reported": reported, "recomputed": recomputed, "abs_diff": abs(reported - recomputed),
                "tol": tol, "verdict": "match" if ok else "DISCREPANCY", "note": note})


def check_set(claim, reported: list, recomputed: list, note=""):
    ok = set(reported) == set(recomputed)
    RES.append({"claim": claim, "reported": ", ".join(reported), "recomputed": ", ".join(recomputed), "abs_diff": np.nan,
                "tol": np.nan, "verdict": "match" if ok else "DISCREPANCY", "note": note})


# ============================================================================ 1. team file
T = load_team()
E = T["emissions"]; IND = T["industries"]
check("team Util/Fun ratio", 6051, E["Util"] / E["Fun"], 1.0)
ties = E.groupby(E.round(9)).apply(lambda s: tuple(sorted(s.index)) if len(s) > 1 else None).dropna()
check("team file tie groups", 3, len(ties), 0, "; ".join(", ".join(t) for t in ties))
rk_team_top = E.rank(ascending=False, method="first")
check("team rank from top: Ships", 2, rk_team_top["Ships"], 0)
check("team rank from top: Aero", 3, rk_team_top["Aero"], 0)

# ============================================================================ 2. mapping (independent)
def sic_to_ff_table():
    rng, cur = {}, None
    for ln in (RAW / "kf_Siccodes49.txt").read_text(errors="ignore").splitlines():
        m_ind = re.match(r"^\s*(\d{1,2})\s+([A-Za-z]+)\b", ln)
        m_rng = re.match(r"^\s*(\d{4})-(\d{4})", ln)
        if m_rng and cur:
            rng[cur].append((int(m_rng.group(1)), int(m_rng.group(2))))
        elif m_ind:
            cur = m_ind.group(2); rng[cur] = []
    look = {}
    overlaps = []
    for ind, rr in rng.items():
        for lo, hi in rr:
            for s in range(lo, hi + 1):
                if s in look and look[s] != ind:
                    overlaps.append((s, look[s], ind))
                look.setdefault(s, ind)
    return look, rng, overlaps


SIC2FF, FFRANGES, OVERLAPS = sic_to_ff_table()
print("FF49 industries parsed:", len(FFRANGES), " overlapping SIC codes across industries:", len(OVERLAPS), OVERLAPS[:5])


def read_pair(fname, a=0, b=2, a_pat=r"\d{6}", b_pat=r"\d{6}"):
    x = pd.read_excel(RAW / fname, header=None, dtype=str)
    A = x.iloc[:, a].astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    B = x.iloc[:, b].astype(str).str.strip().str.replace(r"\.0$", "", regex=True)
    m = A.str.fullmatch(a_pat) & B.str.fullmatch(b_pat)
    return pd.DataFrame({"a": A[m].astype(int), "b": B[m].astype(int)}).drop_duplicates()


s1217 = read_pair("census_2012_to_2017_NAICS.xlsx").rename(columns={"a": "n12", "b": "n17"})
s0712 = read_pair("census_2007_to_2012_NAICS.xls").rename(columns={"a": "n07", "b": "n12"})
s0207 = read_pair("census_2002_to_2007_NAICS.xls").rename(columns={"a": "n02", "b": "n07"})
n02sic = read_pair("census_2002_NAICS_to_1987_SIC.xls", 0, 2, r"\d{6}", r"\d{2,4}").rename(columns={"a": "n02", "b": "sic"})
ch = s1217.merge(s0712, on="n12").merge(s0207, on="n07").merge(n02sic, on="n02")
pairs = ch[["n17", "sic"]].drop_duplicates()
# cross-check against the opposite-direction file the builder used
sic_n02_b = read_pair("census_1987_SIC_to_2002_NAICS.xls", 0, 2, r"\d{2,4}", r"\d{6}").rename(columns={"a": "sic", "b": "n02"})
d1 = set(map(tuple, n02sic[["n02", "sic"]].values)); d2 = set(map(tuple, sic_n02_b[["n02", "sic"]].values))
print("SIC<->NAICS02 link sets: reverse-file", len(d1), " builder-file", len(d2), " symmetric diff", len(d1 ^ d2))

epa = pd.read_csv(RAW / "epa_sc_ghg_naics_v13.csv")
epa.columns = ["naics", "title", "ghg", "unit", "nomargin", "margin", "sef", "ref"]
epa["naics"] = epa["naics"].astype(int); epa["ref"] = epa["ref"].astype(str).str.strip()
unmapped_before_manual = sorted(set(epa["naics"]) - set(pairs["n17"]))
print("EPA codes with no numeric SIC via the chain (before manual links):", unmapped_before_manual)
MANUAL = {112130: 212, 541120: 7389}      # builder's documented manual links, re-applied only if still unmapped
for k, v in MANUAL.items():
    if k not in set(pairs["n17"]):
        pairs = pd.concat([pairs, pd.DataFrame({"n17": [k], "sic": [v]})], ignore_index=True)
pairs["ff"] = pairs["sic"].map(SIC2FF)
asg = pairs.dropna(subset=["ff"])
cnt = asg.groupby(["n17", "ff"])["sic"].nunique().rename("k").reset_index()
cnt["mx"] = cnt.groupby("n17")["k"].transform("max")
prim = cnt[cnt["k"] == cnt["mx"]][["n17", "ff"]]
anyl = cnt[["n17", "ff"]]
OVR = [541713, 541714, 541715]
prim = pd.concat([prim[~prim["n17"].isin(OVR)], pd.DataFrame({"n17": OVR, "ff": "BusSv"})], ignore_index=True)
anyl = pd.concat([anyl, pd.DataFrame({"n17": OVR, "ff": "BusSv"})], ignore_index=True).drop_duplicates()
n_ties = int((prim.groupby("n17").size() > 1).sum())
epa_codes = set(epa["naics"])
check("EPA codes mapped (primary)", 1015, len(epa_codes & set(prim["n17"])), 0)
check("NAICS codes tied under plurality (all NAICS17 incl. electricity)", 115, n_ties, 0,
      "builder counts ties over its mapping table; mine over every NAICS17 in the chain")

# ---- USEEIO v2.0.1 straight from the workbook
import openpyxl
wb = openpyxl.load_workbook(RAW / "USEEIOv2.0.1-411.xlsx", read_only=True)
def ghg_row(sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = [str(c).replace("/US", "") for c in rows[0][1:]]
    r = [row for row in rows[1:] if row[0] == "Greenhouse Gases"][0]
    return pd.Series([float(v) for v in r[1:]], index=hdr)
UN, UD = ghg_row("N"), ghg_row("D")
def ulook(ref, U):
    vals = []
    for c in [c.strip() for c in ref.split(",")]:
        if c in U.index:
            vals.append(U[c]); continue
        for k in range(len(c) - 1, 2, -1):
            hit = [x for x in U.index if x.startswith(c[:k])]
            if hit:
                vals += list(U[hit]); break
    return float(np.mean(vals)) if vals else np.nan
exact = epa["ref"].isin(UN.index)
ratio_exact = (epa.loc[exact, "nomargin"] / epa.loc[exact, "ref"].map(UN)).median()
epa["direct"] = [ulook(r, UD) for r in epa["ref"]]
ratio_all = (epa["nomargin"] / pd.Series([ulook(r_, UN) for r_ in epa["ref"]], index=epa.index)).median()
check("electricity patch, exact-ref median ratio (my method)", 2.642, UN["221100"] * ratio_exact, 0.01, f"USEEIO N(221100)={UN['221100']:.4f}; median ratio exact refs={ratio_exact:.4f}")
patch = UN["221100"] * ratio_all
check("electricity patch, all-ref median ratio with prefix fallback (builder's method)", 2.642, patch, 0.001, f"median ratio all refs={ratio_all:.4f}")
ELEC = [221111, 221112, 221113, 221114, 221115, 221116, 221117, 221118, 221121, 221122]
fac = pd.concat([epa[["naics", "sef", "nomargin", "direct"]],
                 pd.DataFrame({"naics": ELEC, "sef": patch, "nomargin": patch, "direct": UD["221100"]})], ignore_index=True)

def agg(mapping):
    m = mapping.merge(fac, left_on="n17", right_on="naics")
    g = m.groupby("ff")
    return pd.DataFrame({"mean": g["sef"].mean(), "median": g["sef"].median(), "nomargin": g["nomargin"].mean(),
                         "direct": g["direct"].mean(), "n": g["sef"].size()})

A = agg(prim); AL = agg(anyl)
FF49 = list(IND.columns)
A = A.reindex(FF49); AL = AL.reindex(FF49)
print("FF49 industries with no EPA NAICS:", list(A.index[A["mean"].isna()]))
epa_mean = A["mean"]
top = epa_mean.sort_values(ascending=False)
for ind, val in [("Util", 2.242), ("Chems", 0.824), ("Other", 0.752), ("Agric", 0.729), ("Coal", 0.724), ("Banks", 0.068),
                 ("Insur", 0.076), ("Softw", 0.079), ("Hardw", 0.094), ("Smoke", 0.101), ("Chips", 0.102), ("Aero", 0.160), ("Ships", 0.320), ("Oil", 0.377)]:
    check(f"EPA FF49 mean: {ind}", val, epa_mean[ind], 0.0015)
rk_epa_top = epa_mean.rank(ascending=False, method="first")
check("EPA rank from top: Aero", 31, rk_epa_top["Aero"], 0)
check("EPA rank from top: Ships", 20, rk_epa_top["Ships"], 0)
check("EPA rank from top: Oil", 13, rk_epa_top["Oil"], 0)
rk_dir_top = A["direct"].rank(ascending=False, method="first")
check("direct GHG rank from top: Aero", 43, rk_dir_top["Aero"], 0)
check("direct GHG rank from top: Ships", 36, rk_dir_top["Ships"], 0)
um = prim[prim["ff"] == "Util"].merge(epa, left_on="n17", right_on="naics")
check("Util mean without electricity codes", 0.910, um["sef"].mean(), 0.002, "codes: " + " ".join(map(str, um["naics"])))
check("EPA FF49 max/min", 33.2, epa_mean.max() / epa_mean.min(), 0.2)
check("direct FF49 max/min", 531, A["direct"].max() / A["direct"].min(), 5)

def legs(s, n):
    s = s.dropna().sort_values()
    return list(s.index[:n]), list(s.index[::-1][:n])
G5, B5 = legs(epa_mean, 5)
check_set("EPA Green 5", ["Banks", "Insur", "Softw", "Hardw", "Smoke"], G5)
check_set("EPA Brown 5", ["Util", "Chems", "Other", "Agric", "Coal"], B5)
TG5, TB5 = legs(E, 5)
check_set("team Green 5", ["Fun", "RlEst", "Drugs", "Telcm", "Fin"], TG5)
check_set("team Brown 5", ["Util", "Ships", "Aero", "Steel", "BldMt"], TB5)
in41 = E.index
rho, p_rho = stats.spearmanr(epa_mean[in41], E[in41])
check("Spearman team vs EPA (41)", 0.700, rho, 0.005)
check("Spearman p", 3.6e-7, p_rho, 0.2e-7)
rho_d, _ = stats.spearmanr(A["direct"][in41], E[in41])
check("Spearman team vs direct GHG (41)", 0.492, rho_d, 0.01)
check("Spearman EPA median vs EPA mean (49)", 0.970, stats.spearmanr(A["median"], epa_mean)[0], 0.005)
check("Spearman EPA any-link vs EPA mean (49)", 0.991, stats.spearmanr(AL["mean"], epa_mean)[0], 0.005)
G41, B41 = legs(epa_mean[in41], 5)
check_set("EPA Brown 5 within the 41", ["Util", "Chems", "Agric", "Coal", "Food"], B41)

# ============================================================================ 4. Aero / Ships calibration
naics_sef = fac.set_index("naics")["sef"]; naics_dir = fac.set_index("naics")["direct"]
ALT = {"Aero": {"mfg": [336411, 336412, 336413], "transport": [481111, 481112, 481211, 481212, 481219]},
       "Ships": {"mfg": [336611, 336510], "transport": [483111, 483112, 483113, 483114, 483211, 483212]}}
base = [i for i in in41 if i not in ("Aero", "Ships")]
def calib(measure_ff, measure_naics):
    x = np.log(measure_ff[base].values); y = np.log(E[base].values)
    X = np.column_stack([np.ones(len(x)), x]); b = np.linalg.lstsq(X, y, rcond=None)[0]
    e = y - X @ b; df = len(y) - 2; s2 = e @ e / df; XtXi = np.linalg.inv(X.T @ X)
    out = {}
    for ind, alts in ALT.items():
        for w, codes in alts.items():
            x0 = np.array([1, np.log(measure_naics.loc[codes].mean())])
            out[(ind, w)] = (np.log(E[ind]) - x0 @ b) / np.sqrt(s2 * (1 + x0 @ XtXi @ x0))
    return out
cal = calib(epa_mean, naics_sef)
_x = np.log(epa_mean[base].values); _y = np.log(E[base].values); _b = np.polyfit(_x, _y, 1)
for ind, codes in [("Aero", ALT["Aero"]["transport"]), ("Ships", ALT["Ships"]["transport"])]:
    pred = np.exp(np.polyval(_b, np.log(naics_sef.loc[codes].mean())))
    print(f"calibration {ind} transport: predicted team value {pred:.3f} vs actual {E[ind]:.3f} (ratio {E[ind]/pred:.2f}x)")
for (ind, w), rep in {("Aero", "mfg"): 2.88, ("Aero", "transport"): 0.94, ("Ships", "mfg"): 2.52, ("Ships", "transport"): 1.21}.items():
    check(f"calibration t EPA {ind} {w}", rep, cal[(ind, w)], 0.02)
cald = calib(A["direct"], naics_dir)
for (ind, w), rep in {("Aero", "mfg"): 2.98, ("Aero", "transport"): 0.74, ("Ships", "mfg"): 3.18, ("Ships", "transport"): 1.45}.items():
    check(f"calibration t direct {ind} {w}", rep, cald[(ind, w)], 0.03)

# ============================================================================ 5. returns
F5 = load_ff5_mom()
# spot-check the loader against the raw KF text: 2010-01 Mkt-RF
raw = (RAW / "kf_F-F_Research_Data_5_Factors_2x3.csv").read_text().splitlines()
ln = [l for l in raw if l.strip().startswith("201001,")][0]
check("KF FF5 2010-01 Mkt-RF parse (decimal)", float(ln.split(",")[1]) / 100, F5.loc["2010-01-31", "Mkt-RF"], 1e-12)
FAC = ["Mkt-RF", "SMB", "HML", "RMW", "CMA", "UMD"]
print("FF5+UMD last month:", F5.dropna(subset=FAC).index.max().date(), " team industries last month:", IND.index.max().date())


def nw(y, X, L=6):
    d = pd.concat([y.rename("y"), X], axis=1, sort=True).dropna()
    Xm = np.column_stack([np.ones(len(d)), d.drop(columns="y").values]); yv = d["y"].values
    XtXi = np.linalg.inv(Xm.T @ Xm); b = XtXi @ Xm.T @ yv; u = yv - Xm @ b
    xu = Xm * u[:, None]; S = xu.T @ xu
    for l in range(1, L + 1):
        G = xu[l:].T @ xu[:-l]; S += (1 - l / (L + 1)) * (G + G.T)
    se = np.sqrt(np.diag(XtXi @ S @ XtXi)); t = b / se
    return {"n": len(d), "b": b, "t": t, "p": 2 * stats.norm.sf(np.abs(t))}


def gb(green, brown):
    return (IND[green].mean(axis=1) - IND[brown].mean(axis=1)).loc["1970-01-31":"2026-07-31"]


W = {"full": ("1970-01-31", "2026-07-31"), "post2010": ("2010-01-31", "2026-07-31"), "validation": ("2010-01-31", "2022-07-31"),
     "holdout": ("2022-08-31", "2026-07-31"), "covid": ("2020-01-31", "2021-12-31"), "last18": ("2025-02-28", "2026-07-31"),
     "last12": ("2025-08-31", "2026-07-31"), "energy": ("2021-01-31", "2022-12-31"), "pre2010": ("1970-01-31", "2009-12-31")}
nan_check = IND.loc["1970":"2026-07", sorted(set(G5 + B5 + TG5 + TB5 + ["Oil", "Coal", "Chips", "Gold"]))]
print("NaNs in leg industries 1970-2026-07:", int(nan_check.isna().sum().sum()), " min monthly return:", round(nan_check.min().min(), 4))

epa_gb, team_gb = gb(G5, B5), gb(TG5, TB5)
r = nw(epa_gb.loc[W["post2010"][0]:W["post2010"][1]], F5[FAC])
check("PRIMARY n", 199, r["n"], 0)
check("PRIMARY alpha ann", 0.0472, 12 * r["b"][0], 0.0005)
check("PRIMARY t", 1.516, r["t"][0], 0.01)
check("PRIMARY p", 0.129, r["p"][0], 0.002)
check("PRIMARY b_SMB", -0.328, r["b"][2], 0.002); check("PRIMARY t_SMB", -3.37, r["t"][2], 0.02)
check("PRIMARY b_CMA", -0.380, r["b"][5], 0.002); check("PRIMARY t_CMA", -2.01, r["t"][5], 0.02)

def pstats(s):
    s = s.dropna(); w = (1 + s).cumprod()
    return {"ann": 12 * s.mean(), "vol": np.sqrt(12) * s.std(ddof=1), "t": nw(s, pd.DataFrame(index=s.index))["t"][0],
            "mdd": (w / w.cummax() - 1).min(), "n": len(s)}
for nm, s, per, rep in [("EPA", epa_gb, "full", 0.025), ("team", team_gb, "full", -0.005), ("EPA", epa_gb, "post2010", 0.064),
                        ("team", team_gb, "post2010", -0.018), ("EPA", epa_gb, "holdout", 0.096), ("team", team_gb, "holdout", -0.054),
                        ("EPA", epa_gb, "validation", 0.054), ("team", team_gb, "validation", -0.007),
                        ("EPA", epa_gb, "covid", -0.010), ("team", team_gb, "covid", 0.018),
                        ("EPA", epa_gb, "last18", 0.132), ("team", team_gb, "last18", -0.147), ("EPA", epa_gb, "last12", 0.140), ("team", team_gb, "last12", -0.199)]:
    ps = pstats(s.loc[W[per][0]:W[per][1]])
    check(f"raw GB ann mean {nm} {per}", rep, ps["ann"], 0.0006)
ps = pstats(epa_gb.loc[W["post2010"][0]:W["post2010"][1]])
check("EPA post2010 NW t of mean (round 2 text: 1.72)", 1.72, ps["t"], 0.005); check("EPA post2010 max DD", -0.415, ps["mdd"], 0.001)
ps = pstats(epa_gb.loc[W["full"][0]:W["full"][1]])
check("EPA full t of mean", 1.40, ps["t"], 0.01); check("EPA full max DD", -0.712, ps["mdd"], 0.001)

# turnover: drift then re-equal-weight, sum |w - 1/n| per leg, both legs, x12
def tov(names):
    g = 1 + IND[names]; w = g.div(g.sum(axis=1), axis=0)
    return (w - 1 / len(names)).abs().sum(axis=1)
tv = (tov(G5) + tov(B5)).loc[W["post2010"][0]:W["post2010"][1]]
check("EPA post2010 annual turnover (both legs, sum|dw|)", 0.78, 12 * tv.mean(), 0.006)

for per, rep_b in [("last18", 1.20), ("last12", 1.38)]:
    rr = nw(epa_gb.loc[W[per][0]:W[per][1]], F5[FAC])
    check(f"EPA {per} FF5UMD market beta", rep_b, rr["b"][1], 0.01)
rc = nw(epa_gb.loc[W["last12"][0]:W["last12"][1]], F5[["Mkt-RF"]])
check("EPA last12 CAPM alpha", -0.060, 12 * rc["b"][0], 0.001)
rc18 = nw(epa_gb.loc[W["last18"][0]:W["last18"][1]], F5[["Mkt-RF"]])
mkt18 = 12 * F5.loc[W["last18"][0]:W["last18"][1], "Mkt-RF"].mean()
print(f"last18 CAPM: alpha {12*rc18['b'][0]:.4f} beta {rc18['b'][1]:.3f}; market excess {mkt18:.4f}; beta x mkt = {rc18['b'][1]*mkt18:.4f} of GB {12*epa_gb.loc[W['last18'][0]:W['last18'][1]].mean():.4f}")

# ---- 9-variant family (post-2010 and full sample)
e41 = epa_mean[in41]
exO = epa_mean.drop("Other")
FAM = {"epa5": legs(epa_mean, 5), "epa8": legs(epa_mean, 8), "epa_median5": legs(A["median"], 5), "epa_nomargin5": legs(A["nomargin"], 5),
       "epa_anylink5": legs(AL["mean"], 5), "epa41_5": legs(e41, 5), "epa41_8": legs(e41, 8), "epa5_exOther": legs(exO, 5),
       "epa_median8": legs(A["median"], 8)}
def holm_adj(p: pd.Series):
    o = p.sort_values(); m = len(o)
    adj = np.minimum(np.maximum.accumulate([(m - i) * v for i, v in enumerate(o.values)]), 1)
    return pd.Series(adj, index=o.index).reindex(p.index)
fam_rows = []
for per in ("post2010", "full", "holdout"):
    ps_ = {}
    for k, (g, b) in FAM.items():
        rr = nw(gb(g, b).loc[W[per][0]:W[per][1]], F5[FAC])
        ps_[k] = rr["p"][0]
        fam_rows.append({"period": per, "spec": k, "alpha": 12 * rr["b"][0], "t": rr["t"][0], "p": rr["p"][0]})
    hp = holm_adj(pd.Series(ps_))
    for row in fam_rows:
        if row["period"] == per:
            row["holm"] = hp[row["spec"]]
fam = pd.DataFrame(fam_rows)
print(fam.round(3).to_string())
fp = fam[fam.period == "post2010"]
check("post2010 family: # raw p<0.05", 0, int((fp["p"] < 0.05).sum()), 0)
check("post2010 family: min raw p", 0.052, fp["p"].min(), 0.002)
check("post2010 family: min Holm p", 0.471, fp["holm"].min(), 0.005)
ff_ = fam[fam.period == "full"]
check("full family: epa5 Holm p", 0.049, float(ff_.loc[ff_.spec == "epa5", "holm"].iloc[0]), 0.002)
RES.append({"claim": "full-sample family: # raw p<0.05 / # Holm<0.05 (FINDINGS: 'does not hold up across variants')",
            "reported": np.nan, "recomputed": f"{int((ff_['p'] < 0.05).sum())} of 9 / {int((ff_['holm'] < 0.05).sum())} of 9",
            "abs_diff": np.nan, "tol": np.nan, "verdict": "info", "note": "round 1 flagged the old wording (F4); round 2 text says 7 of 9 raw, 5 of 9 Holm"})
check("holdout family: max t (round 2 text: 1.74)", 1.74, fam[fam.period == "holdout"]["t"].max(), 0.005)

# ============================================================================ 6. energy
def ann(s, per): return 12 * s.loc[W[per][0]:W[per][1]].mean()
check("energy 2021-22 team GB ann", -0.192, ann(team_gb, "energy"), 0.001)
check("energy 2021-22 team+Oil+Coal GB ann", -0.348, ann(gb(TG5, TB5 + ["Oil", "Coal"]), "energy"), 0.001)
check("energy 2021-22 EPA GB ann", -0.194, ann(epa_gb, "energy"), 0.001)
check("energy 2021-22 EPA+Oil GB ann", -0.240, ann(gb(G5, B5 + ["Oil"]), "energy"), 0.001)
_, Bx = legs(epa_mean.drop(["Oil", "Coal"]), 5)
check_set("EPA Brown ex Oil/Coal", ["Util", "Chems", "Other", "Agric", "Gold"], Bx)
check("energy 2021-22 EPA ex Oil/Coal GB ann", -0.007, ann(gb(G5, Bx), "energy"), 0.001)
tb = IND[TB5 + ["Oil", "Coal"]].mean(axis=1).loc[W["energy"][0]:W["energy"][1]]
check("energy team+Oil+Coal Brown cum", 0.836, (1 + tb).prod() - 1, 0.002)
cg = IND[["Coal", "Gold"]].loc[W["energy"][0]:W["energy"][1]]
print("2021-22 cumulative: Coal", round(float((1 + cg["Coal"]).prod() - 1), 3), " Gold", round(float((1 + cg["Gold"]).prod() - 1), 3),
      "| EPA GB with Brown minus Coal only (4 industries):", round(ann(gb(G5, [x for x in B5 if x != "Coal"]), "energy"), 4))
nf = load_kf_industries("nfirms")
for ind_, rep in [("Softw", 1), ("Other", 4), ("Coal", 4), ("Agric", 5)]:
    check(f"KF nfirms 1970-01 {ind_}", rep, float(nf.loc["1970-01-31", ind_]), 0)
d_ = (epa_gb - team_gb).loc[W["post2010"][0]:W["post2010"][1]]
check("EPA minus team GB post2010 raw ann (not in FINDINGS text)", 0.0819, 12 * d_.mean(), 0.0005)
check("EPA minus team GB post2010 NW t (not in FINDINGS text)", 2.84, nw(d_, pd.DataFrame(index=d_.index))["t"][0], 0.01)

# ============================================================================ 7. team timing rule, own implementation
ff3t, att = T["ff3"], T["macro"]["attention"]
FF3C = ["Mkt-RF", "SMB", "HML"]
def team_rule(brown, hold, a, b):
    y = (IND[brown].mean(axis=1) - ff3t["RF"])
    d = pd.concat([y.rename("y"), ff3t[FF3C]], axis=1).dropna()
    bet = pd.DataFrame(np.nan, index=d.index, columns=["c"] + FF3C)
    Xa = np.column_stack([np.ones(len(d)), d[FF3C].values]); ya = d["y"].values
    for i in range(59, len(d)):
        bet.iloc[i] = np.linalg.lstsq(Xa[i - 59:i + 1], ya[i - 59:i + 1], rcond=None)[0]
    lag = bet.shift(1)
    eps = d["y"] - (lag[FF3C] * d[FF3C]).sum(axis=1, min_count=3) - lag["c"]
    lz = np.log1p(att); mu = lz.rolling(60, min_periods=36).mean(); sd = lz.rolling(60, min_periods=36).std(ddof=1).replace(0, np.nan)
    z = (lz - mu) / sd
    thr = z.shift(1).expanding(min_periods=60).quantile(0.8)
    st = (z > thr) & thr.notna()
    cross = st & ~st.shift(1, fill_value=False)
    hs = cross.rolling(hold, min_periods=1).max().astype(bool)
    sig = eps.rolling(36, min_periods=36).std(ddof=1)
    mag = (0.05 / (np.sqrt(12) * sig)).clip(upper=1)
    pos = pd.Series(np.nan, index=eps.index)
    live = pos.index >= sig.first_valid_index()
    pos[live] = (-mag).where(hs.reindex(pos.index).fillna(False).astype(bool), 0)[live]
    idx = pos.index.intersection(d.index); idx = idx[idx >= pos.first_valid_index()]
    h = pos.reindex(idx).fillna(0); be = bet[FF3C].reindex(idx); f = ff3t[FF3C].reindex(idx)
    ov = be.mul(-h, axis=0)
    gross = h.shift(1) * d["y"].reindex(idx) + (ov.shift(1) * f).sum(axis=1, min_count=3)
    cost = 10e-4 * h.diff().abs().fillna(h.abs()) + 5e-4 * ov["Mkt-RF"].diff().abs().fillna(ov["Mkt-RF"].abs()) \
        + 25e-4 * ov[["SMB", "HML"]].diff().abs().fillna(ov[["SMB", "HML"]].abs()).sum(axis=1)
    net = (gross - cost.shift(1).fillna(0)).loc[a:b].dropna()
    fit = nw(net, ff3t[FF3C].loc[net.index])
    return 12 * net.mean(), fit["t"][0], 12 * fit["b"][0]
m, t_, _ = team_rule(TB5, 3, *W["holdout"]); check("team rule team Brown hold3 holdout ann_net", -0.0382, m, 0.0005); check("team rule team Brown hold3 holdout alpha t", -2.187, t_, 0.01)
m, t_, _ = team_rule(TB5, 6, *W["holdout"]); check("team rule team Brown hold6 holdout alpha t", -1.657, t_, 0.01)
m, t_, _ = team_rule(TB5, 3, *W["covid"]); check("team rule team Brown hold3 COVID ann_net", 0.059, m, 0.001); check("team rule team Brown hold3 COVID alpha t", 3.21, t_, 0.02)
m, t_, _ = team_rule(TB5, 6, *W["covid"]); check("team rule team Brown hold6 COVID ann_net", 0.074, m, 0.001)
m, t_, _ = team_rule(B5, 3, *W["covid"]); check("team rule EPA Brown hold3 COVID ann_net", 0.005, m, 0.001); check("team rule EPA Brown hold3 COVID alpha t", 0.50, t_, 0.02)
m, t_, _ = team_rule(B5, 6, *W["covid"]); check("team rule EPA Brown hold6 COVID ann_net", 0.008, m, 0.001)
m, t_, _ = team_rule(B5, 3, *W["holdout"]); check("team rule EPA Brown hold3 holdout ann_net", -0.014, m, 0.001)
m, t_, _ = team_rule(["Util", "Steel", "BldMt", "Chems", "Trans"], 3, *W["covid"]); check("team rule team Brown ex Aero/Ships hold3 COVID ann_net", 0.047, m, 0.001)
m, t_, a_ = team_rule(["Util", "Steel", "BldMt", "Chems", "Trans"], 6, *W["covid"]); check("team rule team Brown ex Aero/Ships hold6 COVID ann_net", 0.019, m, 0.001)
m, t_, a_ = team_rule(B5, 3, *W["last12"]); print(f"EPA Brown hold3 last12: ann_net {m:.4f}, FF3 alpha {a_:.4f}, t {t_:.2f}")
m, t_, a_ = team_rule(B5, 6, *W["last18"]); print(f"EPA Brown hold6 last18: ann_net {m:.4f}, FF3 alpha {a_:.4f}, t {t_:.2f}")

# ============================================================================ 8. ledger
L = pd.read_csv(TABLES / "M4_emissions_tests_ledger.csv")
check("ledger rows (round 2: 750 + 32 cal-year + 7 Kendall + 13 sensitivity)", 802, len(L), 0); check("ledger primary rows", 1, int((L.primary_or_exploratory == "primary").sum()), 0)
pr = L[L.primary_or_exploratory == "primary"].iloc[0]
check("ledger primary p vs my p", pr["p_value_two_sided"], r["p"][0], 0.001, pr["test_id"])
AL_ = pd.read_csv(TABLES / "M4_emissions_alphas.csv")
AL_["test_id"] = "alpha_" + AL_["model"] + "_" + AL_["spec"] + "_" + AL_["period"]
mm = L.merge(AL_[["test_id", "p_alpha", "t_alpha"]], on="test_id")
check("ledger alpha rows matched to alphas.csv", 480, len(mm), 0)
check("ledger vs alphas.csv max |p diff|", 0.0, float((mm["p_value_two_sided"] - mm["p_alpha"]).abs().max()), 1e-9)
PF = pd.read_csv(TABLES / "M4_emissions_perf.csv"); PF["test_id"] = "mean_" + PF["spec"] + "_" + PF["period"]
mm2 = L.merge(PF[["test_id", "p_mean_nw"]], on="test_id")
check("ledger mean rows matched to perf.csv (round 2: + 32 cal2021/cal2022)", 192, len(mm2), 0)
check("ledger vs perf.csv max |p diff|", 0.0, float((mm2["p_value_two_sided"] - mm2["p_mean_nw"]).abs().max()), 1e-9)
TR = pd.read_csv(TABLES / "M4_emissions_team_rule_rerun.csv"); TR["test_id"] = "teamrule_" + TR["brown_leg"] + "_" + TR["hold"] + "_" + TR["period"]
mm3 = L.merge(TR[["test_id", "p_alpha"]], on="test_id")
check("ledger team-rule rows matched", 56, len(mm3), 0)
check("ledger vs team_rule_rerun.csv max |p diff|", 0.0, float((mm3["p_value_two_sided"] - mm3["p_alpha"]).abs().max()), 1e-9)
check("ledger p in [0,1] and non-missing", 0, int((~L["p_value_two_sided"].between(0, 1)).sum()), 0)
fam_csv = pd.read_csv(TABLES / "M4_emissions_post2010_family.csv")
fcp = fam_csv[fam_csv.period == "post2010"].set_index("spec")
mine = fam[fam.period == "post2010"].set_index("spec")
check("post2010 family alpha max |diff| vs CSV", 0.0, float((fcp["alpha_ann"] - mine["alpha"]).abs().max()), 0.0005)
check("post2010 family t max |diff| vs CSV", 0.0, float((fcp["t_alpha"] - mine["t"]).abs().max()), 0.01)

# ============================================================================ 9. sensitivity (exploratory, not in FINDINGS)
SENS = []
def fam_alpha(g, b, per):
    rr = nw(gb(g, b).loc[W[per][0]:W[per][1]], F5[FAC]); return 12 * rr["b"][0], rr["t"][0], rr["p"][0], 12 * gb(g, b).loc[W[per][0]:W[per][1]].mean()
for lab, g, b in [("epa5 (primary)", G5, B5), ("Green 5th member Chips instead of Smoke", [x for x in G5 if x != "Smoke"] + ["Chips"], B5),
                  ("Green without Smoke (4 industries)", [x for x in G5 if x != "Smoke"], B5)]:
    for per in ("post2010", "holdout", "full"):
        a, t_, p_, mu = fam_alpha(g, b, per)
        SENS.append({"variant": lab, "period": per, "raw_ann": mu, "ff5umd_alpha": a, "t": t_, "p": p_})
for drop in G5 + B5:
    g = [x for x in G5 if x != drop]; b = [x for x in B5 if x != drop]
    a, t_, p_, mu = fam_alpha(g, b, "post2010")
    SENS.append({"variant": f"drop {drop} (4-industry leg)", "period": "post2010", "raw_ann": mu, "ff5umd_alpha": a, "t": t_, "p": p_})
SENS = pd.DataFrame(SENS)
SENS.to_csv(HERE / "verify_sensitivity.csv", index=False)
print(SENS.round(3).to_string())

# ============================================================================ 10. ROUND 2 (after the builder's fixes F1-F11)
# Independent recomputation of every number the revised FINDINGS text added or changed. Still no module imports.
print("\n==================== ROUND 2 ====================")
from common import load_kf_ff3
FF3K = load_kf_ff3()
def c2(claim, reported, recomputed, tol, note=""):
    check("R2 " + claim, reported, recomputed, tol, note)

# ---- 10a. F3: calibration ratios, direct CO2 measure, Holm/BH over the 12 tests
rowsB = wb["B"].iter_rows(values_only=True)
hdrB = [str(c).replace("/US", "") for c in next(rowsB)[1:]]
co2 = None
for row in rowsB:
    if row[0] == "Carbon dioxide/emission/air/kg":
        co2 = pd.Series([float(v) for v in row[1:]], index=hdrB); break
assert co2 is not None
epa["dco2"] = [ulook(r_, co2) for r_ in epa["ref"]]
fac2 = pd.concat([epa[["naics", "dco2"]], pd.DataFrame({"naics": ELEC, "dco2": co2["221100"]})], ignore_index=True)
DCO2 = prim.merge(fac2, left_on="n17", right_on="naics").groupby("ff")["dco2"].mean().reindex(FF49)
naics_co2 = fac2.set_index("naics")["dco2"]

def calib_full(measure_ff, measure_naics):
    x = np.log(measure_ff[base].values); y = np.log(E[base].values)
    X = np.column_stack([np.ones(len(x)), x]); b = np.linalg.lstsq(X, y, rcond=None)[0]
    e = y - X @ b; df = len(y) - 2; s2 = e @ e / df; XtXi = np.linalg.inv(X.T @ X)
    r2 = 1 - (e @ e) / ((y - y.mean()) @ (y - y.mean()))
    out = {}
    for ind, alts in ALT.items():
        for w, codes in alts.items():
            x0 = np.array([1, np.log(measure_naics.loc[codes].mean())])
            pred = x0 @ b
            t = (np.log(E[ind]) - pred) / np.sqrt(s2 * (1 + x0 @ XtXi @ x0))
            out[(ind, w)] = {"t": t, "p": 2 * stats.t.sf(abs(t), df), "ratio": E[ind] / np.exp(pred), "pred": np.exp(pred)}
    return out, r2

CAL = {}
for nm, mff, mn in [("EPA", epa_mean, naics_sef), ("direct", A["direct"], naics_dir), ("co2", DCO2, naics_co2)]:
    o, r2 = calib_full(mff, mn); CAL[nm] = (o, r2)
ps12 = pd.Series({(nm, ind, w): CAL[nm][0][(ind, w)]["p"] for nm in CAL for ind in ("Aero", "Ships") for w in ("mfg", "transport")})
holm12 = holm_adj(ps12)
def bh_adj(p):
    o = p.sort_values(); m = len(o); q = o.values * m / np.arange(1, m + 1)
    q = np.minimum.accumulate(q[::-1])[::-1]
    return pd.Series(np.minimum(q, 1), index=o.index).reindex(p.index)
bh12 = bh_adj(ps12)
REP_RATIO = {("EPA", "Aero", "mfg"): 21.4, ("EPA", "Aero", "transport"): 2.8, ("EPA", "Ships", "mfg"): 14.5, ("EPA", "Ships", "transport"): 3.8,
             ("direct", "Aero", "mfg"): 38.8, ("direct", "Aero", "transport"): 2.5, ("direct", "Ships", "mfg"): 47.0, ("direct", "Ships", "transport"): 6.0,
             ("co2", "Aero", "mfg"): 36.9, ("co2", "Aero", "transport"): 1.9, ("co2", "Ships", "mfg"): 42.4, ("co2", "Ships", "transport"): 4.9}
REP_P = {("EPA", "Aero", "mfg"): 0.007, ("EPA", "Aero", "transport"): 0.355, ("EPA", "Ships", "mfg"): 0.016, ("EPA", "Ships", "transport"): 0.234,
         ("direct", "Aero", "mfg"): 0.005, ("direct", "Aero", "transport"): 0.463, ("direct", "Ships", "mfg"): 0.003, ("direct", "Ships", "transport"): 0.157,
         ("co2", "Aero", "mfg"): 0.005, ("co2", "Aero", "transport"): 0.623, ("co2", "Ships", "mfg"): 0.003, ("co2", "Ships", "transport"): 0.202}
for k, v in REP_RATIO.items():
    c2(f"calibration team/predicted {k}", v, CAL[k[0]][0][(k[1], k[2])]["ratio"], 0.051)
for k, v in REP_P.items():
    c2(f"calibration raw p (t, 37 df) {k}", v, ps12[k], 0.0006)
for k, v in {("co2", "Aero", "mfg"): 2.99, ("co2", "Aero", "transport"): 0.50, ("co2", "Ships", "mfg"): 3.14, ("co2", "Ships", "transport"): 1.30}.items():
    c2(f"calibration t {k}", v, CAL["co2"][0][(k[1], k[2])]["t"], 0.006)
for k, v in {("EPA", "Aero", "mfg"): 0.052, ("EPA", "Ships", "mfg"): 0.114, ("direct", "Aero", "mfg"): 0.050, ("direct", "Ships", "mfg"): 0.036,
             ("co2", "Aero", "mfg"): 0.050, ("co2", "Ships", "mfg"): 0.037}.items():
    c2(f"calibration Holm p over 12 {k}", v, holm12[k], 0.0006)
for k, v in {("EPA", "Aero", "mfg"): 0.016, ("EPA", "Ships", "mfg"): 0.033}.items():
    c2(f"calibration BH p over 12 {k}", v, bh12[k], 0.0006)
c2("calibration EPA transport predicted Aero", 0.192, CAL["EPA"][0][("Aero", "transport")]["pred"], 0.0006)
c2("calibration EPA transport predicted Ships", 0.271, CAL["EPA"][0][("Ships", "transport")]["pred"], 0.0006)
c2("calibration fit R2 EPA", 0.54, CAL["EPA"][1], 0.005)
c2("calibration fit R2 direct GHG", 0.40, CAL["direct"][1], 0.005)
c2("calibration fit R2 direct CO2", 0.42, CAL["co2"][1], 0.005)
RES.append({"claim": "R2 calibration: # of 12 with Holm p<0.05 (text: the 4 direct-measure mfg rejections survive, EPA ones do not)",
            "reported": np.nan, "recomputed": ", ".join("/".join(k) for k, v in holm12.items() if v < 0.05), "abs_diff": np.nan, "tol": np.nan,
            "verdict": "match" if set(k for k, v in holm12.items() if v < 0.05) == {("direct", "Aero", "mfg"), ("direct", "Ships", "mfg"), ("co2", "Aero", "mfg"), ("co2", "Ships", "mfg")} else "DISCREPANCY", "note": ""})
cc = pd.read_csv(TABLES / "M4_emissions_aero_ships_calibration.csv")
mp = {"epa_sc_mean": "EPA", "useeio_direct_mean": "direct", "useeio_direct_co2_mean": "co2"}
cc["key"] = list(zip(cc["measure"].map(mp), cc["industry"], cc["mapping"]))
c2("calibration CSV holm_p_12 vs mine, max |diff|", 0.0, float((cc["holm_p_12"] - cc["key"].map(holm12)).abs().max()), 1e-6)
c2("calibration CSV bh_p_12 vs mine, max |diff|", 0.0, float((cc["bh_p_12"] - cc["key"].map(bh12)).abs().max()), 1e-6)
c2("calibration CSV ratio vs mine, max |diff|", 0.0, float((cc["team_over_predicted"] - cc["key"].map(lambda k: CAL[k[0]][0][(k[1], k[2])]["ratio"])).abs().max()), 1e-6)

# ---- 10b. F2: rank evidence and spreads
UT = prim.merge(pd.concat([epa[["naics"]].assign(tot=[ulook(r_, UN) for r_ in epa["ref"]]),
                           pd.DataFrame({"naics": ELEC, "tot": UN["221100"]})], ignore_index=True),
                left_on="n17", right_on="naics").groupby("ff")["tot"].mean().reindex(FF49)
c2("Spearman team vs USEEIO total (41)", 0.649, stats.spearmanr(UT[in41], E[in41])[0], 0.005)
c2("Spearman team vs USEEIO direct CO2 (41)", 0.503, stats.spearmanr(DCO2[in41], E[in41])[0], 0.005)
c2("USEEIO total FF49 max/min", 45.2, UT.max() / UT.min(), 0.3)
c2("EPA NAICS max/min (with margins)", 135.3, epa["sef"].max() / epa["sef"].min(), 0.5)
c2("EPA NAICS min / electricity factor", 0.011, epa["sef"].min() / patch, 0.0006)
kt = stats.kendalltau(epa_mean[in41], E[in41])
c2("Kendall tau team vs EPA (41)", 0.507, kt[0], 0.0006); c2("Kendall p", 3.3e-6, kt[1], 0.06e-6)
c2("Spearman EPA nomargin vs EPA mean (49)", 0.987, stats.spearmanr(A["nomargin"], epa_mean)[0], 0.0006)
rr_all = (epa["nomargin"] / pd.Series([ulook(r_, UN) for r_ in epa["ref"]], index=epa.index))
c2("electricity scale ratio IQR low", 0.459, rr_all.quantile(0.25), 0.0006); c2("electricity scale ratio IQR high", 0.690, rr_all.quantile(0.75), 0.0006)
G8, B8 = legs(epa_mean, 8); TG8, TB8 = legs(E, 8)
c2("8-per-leg Green overlap EPA vs team", 4, len(set(G8) & set(TG8)), 0); c2("8-per-leg Brown overlap", 3, len(set(B8) & set(TB8)), 0)
srt = epa_mean.sort_values()
c2("Smoke EPA mean (4 dp)", 0.1010, srt["Smoke"], 0.00006); c2("Chips EPA mean (4 dp)", 0.1022, srt["Chips"], 0.00006)
c2("Green 5th/6th margin (Chips - Smoke)", 0.0012, srt.iloc[5] - srt.iloc[4], 0.00006)
c2("Brown 5th/6th margin (Coal - Gold)", 0.131, srt.iloc[-5] - srt.iloc[-6], 0.0006)
for ind_, rep in [("Chems", 2), ("Coal", 5), ("Trans", 8), ("Oil", 13), ("Mines", 17)]:
    c2(f"EPA rank from top {ind_}", rep, rk_epa_top[ind_], 0)
for ind_, rep in [("Chems", 6), ("Coal", 9), ("Trans", 7), ("Mines", 8)]:
    c2(f"team rank from top {ind_}", rep, rk_team_top[ind_], 0)
for ind_, rep in [("Chems", 6), ("Coal", 2), ("Trans", 7), ("Oil", 9), ("Mines", 8)]:
    c2(f"direct GHG rank from top {ind_}", rep, rk_dir_top[ind_], 0)
for ind_, rep in [("Aero", 0.531), ("Ships", 1.019), ("Steel", 0.403), ("BldMt", 0.365)]:
    c2(f"team value {ind_}", rep, E[ind_], 0.0006)
air = [c for c in epa["naics"] if str(c).startswith("481")]; wat = [c for c in epa["naics"] if str(c).startswith("483")]
c2("EPA air transport 481 mean", 0.644, naics_sef.loc[air].mean(), 0.0006); c2("EPA water 483 mean", 0.816, naics_sef.loc[wat].mean(), 0.0006)
c2("direct air 481 mean", 0.726, naics_dir.loc[air].mean(), 0.0006); c2("direct water 483 mean", 0.516, naics_dir.loc[wat].mean(), 0.0006)
c2("EPA aircraft mfg min (336411-3)", 0.139, naics_sef.loc[[336411, 336412, 336413]].min(), 0.0006)
c2("EPA aircraft mfg max (336411-3)", 0.170, naics_sef.loc[[336411, 336412, 336413]].max(), 0.0006)
c2("EPA ship building 336611", 0.196, naics_sef.loc[336611], 0.0006); c2("direct ship building 336611", 0.017, naics_dir.loc[336611], 0.0006)
c2("direct aircraft mfg min", 0.005, naics_dir.loc[[336411, 336412, 336413]].min(), 0.0006)
c2("direct aircraft mfg max", 0.013, naics_dir.loc[[336411, 336412, 336413]].max(), 0.0006)

for ind_, rep in [("Gold", 0.593), ("Food", 0.524), ("Trans", 0.501), ("Steel", 0.448), ("Fin", 0.104), ("Fun", 0.110)]:
    c2(f"EPA FF49 mean {ind_}", rep, epa_mean[ind_], 0.0006)
c2("Util n_naics", 13, A.loc["Util", "n"], 0); c2("Smoke n_naics", 1, A.loc["Smoke", "n"], 0)
c2("holdout family max t is without margins", 1, int(fam[fam.period == "holdout"].sort_values("t").iloc[-1]["spec"] == "epa_nomargin5"), 0)

# ---- 10c. F8: leg-membership sensitivity vs the new table
PS = pd.read_csv(TABLES / "M4_emissions_primary_sensitivity.csv")
def sens_row(g, b, per):
    a, t_, p_, mu = fam_alpha(g, b, per); return a, t_, p_
ch_rows = PS[PS.variant.str.startswith("Chips")]
for _, rw in ch_rows.iterrows():
    per = {"full_1970": "full"}.get(rw["period"], rw["period"])
    a, t_, p_ = sens_row([x for x in G5 if x != "Smoke"] + ["Chips"], B5, per)
    c2(f"sens Chips-for-Smoke {per} t (CSV vs mine)", rw["t_alpha"], t_, 1e-6); c2(f"sens Chips-for-Smoke {per} alpha", rw["alpha_ann"], a, 1e-8)
for _, rw in PS[PS.kind == "drop-one"].iterrows():
    drop = rw["variant"].split()[1]
    a, t_, p_ = sens_row([x for x in G5 if x != drop], [x for x in B5 if x != drop], "post2010")
    c2(f"sens drop {drop} post2010 p (CSV vs mine)", rw["p_alpha"], p_, 1e-6)
d1 = PS[PS.kind == "drop-one"]
c2("drop-one alpha min (text 3.1%)", 0.031, d1["alpha_ann"].min(), 0.0006); c2("drop-one alpha max (text 5.6%)", 0.056, d1["alpha_ann"].max(), 0.0006)
c2("drop-one p min (text 0.087, Coal)", 0.087, d1["p_alpha"].min(), 0.0006); c2("drop-one p max (text 0.264, Chems)", 0.264, d1["p_alpha"].max(), 0.0006)
nfS = nf["Smoke"]
c2("Smoke nfirms 1970-01", 10, float(nfS.loc["1970-01-31"]), 0); c2("Smoke nfirms 2010-01", 6, float(nfS.loc["2010-01-31"]), 0)
c2("Smoke nfirms 2026-07", 5, float(nfS.loc["2026-07-31"]), 0); c2("Smoke nfirms min 1970-2026", 3, float(nfS.loc["1970":"2026-07"].min()), 0)

# ---- 10d. F11: recent-window CAPM decomposition (KF FF3-file market, as the module states)
RB = pd.read_csv(TABLES / "M4_emissions_recent_beta.csv")
SP = {"epa5": epa_gb, "team5": team_gb, "epa5_minus_team5": epa_gb - team_gb}
worst = 0.0
for _, rw in RB.iterrows():
    s = SP[rw["spec"]].loc[W[rw["period"]][0]:W[rw["period"]][1]]
    rc_ = nw(s, FF3K[["Mkt-RF"]]); mk = 12 * FF3K["Mkt-RF"].reindex(s.index).mean()
    for col, val in [("capm_beta", rc_["b"][1]), ("capm_alpha_ann", 12 * rc_["b"][0]), ("capm_beta_x_mkt", rc_["b"][1] * mk), ("ann_ret", 12 * s.mean()), ("ann_mkt_excess", mk)]:
        worst = max(worst, abs(rw[col] - val))
    c2(f"identity ann_ret = alpha + beta x mkt ({rw['spec']} {rw['period']})", 12 * s.mean(), 12 * rc_["b"][0] + rc_["b"][1] * mk, 1e-10)
c2("recent_beta.csv vs my CAPM, max |diff| over all cells", 0.0, worst, 1e-8)
def capm(s, per):
    s = s.loc[W[per][0]:W[per][1]]; rc_ = nw(s, FF3K[["Mkt-RF"]]); mk = 12 * FF3K["Mkt-RF"].reindex(s.index).mean()
    return rc_["b"][1], rc_["b"][1] * mk, 12 * rc_["b"][0], mk, 12 * s.mean()
b_, bx, al, mk, mu = capm(epa_gb, "last18")
c2("last18 CAPM beta", 0.91, b_, 0.006); c2("last18 mkt excess", 0.119, mk, 0.0006); c2("last18 beta x mkt", 0.108, bx, 0.0006); c2("last18 CAPM alpha", 0.024, al, 0.0006)
b_, bx, al, mk, mu = capm(epa_gb, "last12")
c2("last12 CAPM beta", 1.37, b_, 0.006); c2("last12 beta x mkt", 0.199, bx, 0.0006); c2("last12 CAPM alpha", -0.060, al, 0.0006)
b_, bx, al, mk, mu = capm(epa_gb, "post2010"); c2("post2010 CAPM beta", 0.06, b_, 0.006); c2("post2010 beta x mkt", 0.008, bx, 0.0006)
b_, bx, al, mk, mu = capm(epa_gb, "holdout"); c2("holdout CAPM beta", 0.26, b_, 0.006); c2("holdout beta x mkt", 0.034, bx, 0.0006)
for per, rb, rt in [("last18", 1.20, 3.72), ("last12", 1.38, 7.01)]:
    rr = nw(epa_gb.loc[W[per][0]:W[per][1]], F5[FAC]); c2(f"{per} FF5UMD mkt beta t", rt, rr["t"][1], 0.006)

# ---- 10e. F7: EPA minus team
dser = epa_gb - team_gb
REPD = {"post2010": (0.082, 2.84, 0.004, (0.064, 2.30), (0.055, 2.08), (0.049, 1.93)),
        "holdout": (0.150, 2.31, 0.021, (0.102, 2.30), (0.096, 2.03), (0.083, 1.45)),
        "validation": (0.060, 2.01, 0.045, (0.051, 1.59), (0.048, 1.52), (0.035, 1.17)),
        "full": (0.030, 1.86, 0.063, (0.020, 1.24), (0.019, 1.18), (0.023, 1.44))}
def three(s, per):
    s = s.loc[W[per][0]:W[per][1]]
    m0 = nw(s, pd.DataFrame(index=s.index))
    out = [(12 * s.mean(), m0["t"][0], m0["p"][0])]
    for Fm in (FF3K[["Mkt-RF"]], FF3K[["Mkt-RF", "SMB", "HML"]], F5[FAC]):
        rr = nw(s, Fm); out.append((12 * rr["b"][0], rr["t"][0], rr))
    return out
for per, (rm, rt, rp, cap, f3, f5) in REPD.items():
    o = three(dser, per)
    c2(f"EPA-team {per} raw mean", rm, o[0][0], 0.0006); c2(f"EPA-team {per} raw t", rt, o[0][1], 0.006); c2(f"EPA-team {per} raw p", rp, o[0][2], 0.0006)
    for lab, rep, got in [("CAPM", cap, o[1]), ("FF3", f3, o[2]), ("FF5UMD", f5, o[3])]:
        c2(f"EPA-team {per} {lab} alpha", rep[0], got[0], 0.0006); c2(f"EPA-team {per} {lab} t", rep[1], got[1], 0.006)

# multiplicity context for the EPA-minus-team table (16 tests shown in FINDINGS; 42 EPA-minus-team rows in the ledger)
p16 = pd.Series({(per, j): three(dser, per)[j][2]["p"][0] if j else three(dser, per)[0][2] for per in REPD for j in range(4)})
fam42 = L[L.test_id.str.contains("epa5_minus_team5")]["p_value_two_sided"]
RES.append({"claim": "R2 EPA-minus-team: min Holm p over the 16 tests in the FINDINGS table / BH / over the 42 ledger rows (post2010 raw row)",
            "reported": np.nan, "recomputed": f"Holm16 {holm_adj(p16).min():.3f}; BH16 {bh_adj(p16).min():.3f}; n42 {len(fam42)}; post2010-raw Holm42 {float(holm_adj(fam42.reset_index(drop=True))[fam42.reset_index(drop=True).sub(p16[('post2010', 0)]).abs().idxmin()]):.3f}",
            "abs_diff": np.nan, "tol": np.nan, "verdict": "info", "note": "FINDINGS says '1 of 802 ledger tests, not adjusted'; see Round 2 N2"})

# ---- 10f. F4: full-sample family and its decomposition; alpha grid; loadings
o = three(epa_gb, "full")
c2("EPA full CAPM alpha", 0.014, o[1][0], 0.0006); c2("EPA full CAPM t", 0.78, o[1][1], 0.006)
c2("EPA full FF3 alpha", 0.024, o[2][0], 0.0006); c2("EPA full FF3 t", 1.43, o[2][1], 0.006)
c2("EPA full FF5UMD alpha", 0.045, o[3][0], 0.0006); c2("EPA full FF5UMD t", 2.58, o[3][1], 0.006)
rfull = o[3][2]
c2("EPA full CMA loading", -0.345, rfull["b"][5], 0.0006); c2("EPA full CMA t", -3.04, rfull["t"][5], 0.006)
c2("EPA full UMD loading", -0.096, rfull["b"][6], 0.0006); c2("EPA full UMD t", -2.02, rfull["t"][6], 0.006)
ffam = fam[fam.period == "full"]
c2("full family # raw p<0.05", 7, int((ffam["p"] < 0.05).sum()), 0); c2("full family # Holm p<0.05", 5, int((ffam["holm"] < 0.05).sum()), 0)
check_set("R2 full family Holm survivors", ["epa5", "epa8", "epa_nomargin5", "epa_anylink5", "epa_median8"], list(ffam.loc[ffam["holm"] < 0.05, "spec"]))
for sp, rep in [("epa8", 0.014), ("epa_nomargin5", 0.011), ("epa_anylink5", 0.046), ("epa_median8", 0.040)]:
    c2(f"full family Holm p {sp}", rep, float(ffam.loc[ffam.spec == sp, "holm"].iloc[0]), 0.0006)
c2("full family median5 t", 1.72, float(ffam.loc[ffam.spec == "epa_median5", "t"].iloc[0]), 0.006)
c2("full family exOther t", 1.48, float(ffam.loc[ffam.spec == "epa5_exOther", "t"].iloc[0]), 0.006)
rpre = nw(epa_gb.loc[W["pre2010"][0]:W["pre2010"][1]], F5[FAC])
c2("pre2010 FF5UMD alpha", 0.040, 12 * rpre["b"][0], 0.0006); c2("pre2010 t", 1.81, rpre["t"][0], 0.006); c2("pre2010 p", 0.070, rpre["p"][0], 0.0006)
pp = pstats(epa_gb.loc[W["pre2010"][0]:W["pre2010"][1]]); c2("pre2010 raw mean", 0.009, pp["ann"], 0.0006); c2("pre2010 raw t", 0.46, pp["t"], 0.006)
pf = fam[fam.period == "post2010"]
c2("post2010 family min alpha (median5)", 0.019, pf["alpha"].min(), 0.0006); c2("post2010 family max alpha (anylink)", 0.055, pf["alpha"].max(), 0.0006)
c2("post2010 exOther alpha", 0.036, float(pf.loc[pf.spec == "epa5_exOther", "alpha"].iloc[0]), 0.0006)
o = three(epa_gb, "post2010")
c2("EPA post2010 CAPM alpha", 0.056, o[1][0], 0.0006); c2("EPA post2010 CAPM t", 1.43, o[1][1], 0.006)
c2("EPA post2010 FF3 alpha", 0.044, o[2][0], 0.0006); c2("EPA post2010 FF3 t", 1.32, o[2][1], 0.006)
rp10 = o[3][2]
for j, nmj, rep in [(1, "Mkt", 0.108), (3, "HML", -0.024), (4, "RMW", -0.035), (6, "UMD", 0.066)]:
    c2(f"primary loading {nmj}", rep, rp10["b"][j], 0.0006)
c2("primary HML t", -0.24, rp10["t"][3], 0.006)
dd = pd.concat([epa_gb.rename("y"), F5[FAC]], axis=1, sort=True).loc[W["post2010"][0]:W["post2010"][1]].dropna()
Xr = np.column_stack([np.ones(len(dd)), dd[FAC].values]); br = np.linalg.lstsq(Xr, dd["y"].values, rcond=None)[0]; er = dd["y"].values - Xr @ br
c2("primary R2", 0.148, 1 - er @ er / ((dd["y"] - dd["y"].mean()) ** 2).sum(), 0.0006)
r3e = nw(epa_gb.loc[W["post2010"][0]:W["post2010"][1]], FF3K[["Mkt-RF", "SMB", "HML"]]); r3t = nw(team_gb.loc[W["post2010"][0]:W["post2010"][1]], FF3K[["Mkt-RF", "SMB", "HML"]])
c2("post2010 FF3 HML EPA", -0.267, r3e["b"][3], 0.0006); c2("post2010 FF3 HML t EPA", -4.12, r3e["t"][3], 0.006)
c2("post2010 FF3 HML team", -0.298, r3t["b"][3], 0.0006); c2("post2010 FF3 HML t team", -5.75, r3t["t"][3], 0.006)
o = three(epa_gb, "holdout")
c2("holdout CAPM alpha", 0.062, o[1][0], 0.0006); c2("holdout CAPM t", 1.09, o[1][1], 0.006)
c2("holdout FF3 alpha", 0.050, o[2][0], 0.0006); c2("holdout FF3 t", 1.05, o[2][1], 0.006)
c2("holdout FF5UMD alpha", 0.039, o[3][0], 0.0006); c2("holdout FF5UMD t", 0.70, o[3][1], 0.006)
for per, ra, rt in [("post2010", -0.002, -0.07), ("full", 0.022, 1.60), ("holdout", -0.044, -0.90)]:
    rr = nw(team_gb.loc[W[per][0]:W[per][1]], F5[FAC]); c2(f"team FF5UMD alpha {per}", ra, 12 * rr["b"][0], 0.0006); c2(f"team FF5UMD t {per}", rt, rr["t"][0], 0.006)

# ---- 10g. performance table (t, vol, Sharpe, DD), turnover and cost
W["inflation_rates"] = ("2022-01-31", "2024-12-31")
def prow(s, per):
    p_ = pstats(s.loc[W[per][0]:W[per][1]]); p_["sharpe"] = p_["ann"] / p_["vol"]
    p_["p"] = 2 * stats.norm.sf(abs(p_["t"])); return p_
for nm, s, per, key, rep, tol in [("EPA", epa_gb, "full", "vol", 0.127, 6e-4), ("EPA", epa_gb, "full", "sharpe", 0.20, 6e-3),
                                  ("team", team_gb, "full", "vol", 0.097, 6e-4), ("team", team_gb, "full", "sharpe", -0.05, 6e-3), ("team", team_gb, "full", "mdd", -0.698, 6e-4), ("team", team_gb, "full", "t", -0.35, 6e-3),
                                  ("EPA", epa_gb, "post2010", "vol", 0.123, 6e-4), ("EPA", epa_gb, "post2010", "sharpe", 0.52, 6e-3), ("EPA", epa_gb, "post2010", "p", 0.085, 6e-4),
                                  ("team", team_gb, "post2010", "vol", 0.099, 6e-4), ("team", team_gb, "post2010", "sharpe", -0.18, 6e-3), ("team", team_gb, "post2010", "mdd", -0.509, 6e-4), ("team", team_gb, "post2010", "t", -0.68, 6e-3),
                                  ("EPA", epa_gb, "validation", "t", 1.22, 6e-3), ("team", team_gb, "validation", "t", -0.22, 6e-3),
                                  ("EPA", epa_gb, "holdout", "t", 1.61, 6e-3), ("team", team_gb, "holdout", "t", -0.99, 6e-3), ("EPA", epa_gb, "holdout", "sharpe", 0.70, 6e-3),
                                  ("EPA", epa_gb, "inflation_rates", "ann", -0.035, 6e-4), ("team", team_gb, "inflation_rates", "ann", -0.076, 6e-4),
                                  ("EPA", epa_gb, "last18", "t", 0.96, 6e-3), ("team", team_gb, "last18", "t", -1.83, 6e-3),
                                  ("EPA", epa_gb, "last12", "t", 0.70, 6e-3), ("team", team_gb, "last12", "t", -2.19, 6e-3)]:
    c2(f"perf {nm} {per} {key}", rep, prow(s, per)[key], tol)
tvt = (tov(TG5) + tov(TB5)).loc[W["post2010"][0]:W["post2010"][1]]
c2("team post2010 turnover (two-way, both legs)", 0.69, 12 * tvt.mean(), 0.006)
c2("EPA post2010 cost at 10bp", 0.0008, 10e-4 * 12 * tv.mean(), 0.00006); c2("team post2010 cost at 10bp", 0.0007, 10e-4 * 12 * tvt.mean(), 0.00006)

# ---- 10h. energy table
def cum(s, per): s = s.loc[W[per][0]:W[per][1]]; return (1 + s).prod() - 1
tpoc = gb(TG5, TB5 + ["Oil", "Coal"]); epo = gb(G5, B5 + ["Oil"]); exoc = gb(G5, Bx)
for nm, s, rep in [("team", team_gb, -0.334), ("team+Oil+Coal", tpoc, -0.517), ("EPA", epa_gb, -0.336), ("EPA+Oil", epo, -0.395), ("EPA exOC", exoc, -0.027)]:
    c2(f"energy cumulative GB {nm}", rep, cum(s, "energy"), 0.0006)
rr = nw(tpoc.loc[W["energy"][0]:W["energy"][1]], F5[FAC]); c2("energy team+OC FF5UMD alpha", -0.269, 12 * rr["b"][0], 0.0006); c2("energy team+OC FF5UMD t", -2.63, rr["t"][0], 0.006)
for nm, legn, rep in [("team", TB5, 0.350), ("EPA", B5, 0.679), ("EPA exOC", Bx, 0.170)]:
    c2(f"energy Brown-leg cumulative {nm}", rep, (1 + IND[legn].mean(axis=1).loc[W["energy"][0]:W["energy"][1]]).prod() - 1, 0.0006)
for nm, s, rep in [("team+OC", tpoc, -0.040), ("EPA+Oil", epo, 0.090), ("EPA exOC", exoc, 0.069)]:
    c2(f"holdout GB ann {nm}", rep, ann(s, "holdout"), 0.0006)

# ---- 10i. F5/F6: timing rule claims
TRW = ["post2010", "validation", "holdout", "covid", "inflation_rates"]
tr = {(h, per): team_rule(B5, h, *W[per]) for h in (3, 6) for per in TRW + ["last12", "last18"]}
mx = max(((abs(v[1]), k) for k, v in tr.items() if k[1] in TRW))
c2("EPA Brown max |FF3 alpha t| over post2010/validation/holdout/covid/inflation_rates", 1.75, mx[0], 0.006, f"at {mx[1]}")
c2("EPA Brown hold3 last12 FF3 t", 2.72, tr[(3, "last12")][1], 0.006); c2("EPA Brown hold3 last12 ann_net", -0.022, tr[(3, "last12")][0], 0.0006)
c2("EPA Brown hold6 last18 FF3 t", -2.77, tr[(6, "last18")][1], 0.006)
c2("EPA Brown hold3 post2010 ann_net", 0.009, tr[(3, "post2010")][0], 0.0006); c2("EPA Brown hold3 post2010 t", 1.13, tr[(3, "post2010")][1], 0.006)
c2("EPA Brown hold6 post2010 ann_net", 0.024, tr[(6, "post2010")][0], 0.0006); c2("EPA Brown hold6 post2010 t", 1.46, tr[(6, "post2010")][1], 0.006)
c2("EPA Brown hold3 holdout t", -1.30, tr[(3, "holdout")][1], 0.006)
c2("EPA Brown hold6 holdout ann_net", 0.001, tr[(6, "holdout")][0], 0.0006); c2("EPA Brown hold6 holdout t", -0.61, tr[(6, "holdout")][1], 0.006)
c2("EPA Brown hold6 COVID t", 0.52, tr[(6, "covid")][1], 0.006)
XAS = ["Util", "Steel", "BldMt", "Chems", "Trans"]
_, rk = legs(E.drop(["Aero", "Ships"]), 5); check_set("R2 team Brown re-ranked without Aero/Ships", XAS, rk)
m, t_, a_ = team_rule(TB5, 6, *W["covid"]); c2("team Brown hold6 COVID FF3 t", 2.07, t_, 0.006); c2("team Brown hold6 COVID FF3 alpha", 0.052, a_, 0.0006)
m, t_, a_ = team_rule(TB5, 6, *W["holdout"]); c2("team Brown hold6 holdout ann_net", -0.021, m, 0.0006)
m, t_, a_ = team_rule(TB5, 3, *W["post2010"]); c2("team Brown hold3 2010-2026 ann_net", 0.0007, m, 0.00006)
m, t_, a_ = team_rule(XAS, 3, *W["covid"]); c2("ex-Aero/Ships hold3 COVID t", 1.47, t_, 0.006)
m, t_, a_ = team_rule(XAS, 6, *W["covid"]); c2("ex-Aero/Ships hold6 COVID FF3 alpha", -0.021, a_, 0.0006); c2("ex-Aero/Ships hold6 COVID t", -0.39, t_, 0.006)
m, t_, a_ = team_rule(XAS, 3, *W["holdout"]); c2("ex-Aero/Ships hold3 holdout ann_net", -0.018, m, 0.0006); c2("ex-Aero/Ships hold3 holdout t", -1.30, t_, 0.006)
m, t_, a_ = team_rule(XAS, 6, *W["holdout"]); c2("ex-Aero/Ships hold6 holdout ann_net", 0.003, m, 0.0006); c2("ex-Aero/Ships hold6 holdout t", 0.10, t_, 0.006)

# ---- 10j. F10: every p-value printed in any M4 table appears in the ledger
Lp = np.sort(L["p_value_two_sided"].values)
def in_ledger(v):
    i = np.searchsorted(Lp, v); return any(0 <= j < len(Lp) and abs(Lp[j] - v) <= 1e-12 * max(1, abs(v)) + 1e-15 for j in (i - 1, i))
miss = {}
for f in sorted(TABLES.glob("M4_emissions_*.csv")):
    if f.name.endswith("tests_ledger.csv"): continue
    T_ = pd.read_csv(f)
    pcols = [c for c in T_.columns if re.fullmatch(r"(p|p_.*|.*_p|.*_p_.*)", c) and not re.search(r"holm|bh|share|pos|pct|n_params", c)]
    for c in pcols:
        v = pd.to_numeric(T_[c], errors="coerce").dropna()
        v = v[(v >= 0) & (v <= 1)]
        k = int(sum(not in_ledger(x) for x in v.values))
        if k: miss[f"{f.name}:{c}"] = k
c2("p-values in M4 tables missing from ledger (unadjusted p columns)", 0, sum(miss.values()), 0, "; ".join(f"{k}={v}" for k, v in miss.items()) or "none")
nk = L[L.test_id.str.startswith("rank_kendall")]; c2("ledger Kendall rows", 7, len(nk), 0)
ncal = L[L.test_id.str.contains(r"_cal202[12]$")]; c2("ledger cal-year rows", 32, len(ncal), 0)
nsens = L[L.test_id.str.startswith("sens_")]; c2("ledger sensitivity rows", 13, len(nsens), 0)
c2("ledger sensitivity rows labelled robustness", 13, int((nsens.primary_or_exploratory == "robustness").sum()), 0)
cnt_lab = L.primary_or_exploratory.value_counts()
c2("ledger robustness rows", 375, int(cnt_lab.get("robustness", 0)), 0); c2("ledger exploratory rows", 426, int(cnt_lab.get("exploratory", 0)), 0)
# my own p for two of the new cal-year rows
for tid, g_, b_ in [("mean_team5_cal2022", TG5, TB5), ("mean_epa5_cal2022", G5, B5)]:
    s = gb(g_, b_).loc["2022-01-31":"2022-12-31"]; pm = 2 * stats.norm.sf(abs(nw(s, pd.DataFrame(index=s.index))["t"][0]))
    c2(f"ledger {tid} p vs mine", float(L.loc[L.test_id == tid, "p_value_two_sided"].iloc[0]), pm, 1e-9)

out = pd.DataFrame(RES)
out.to_csv(HERE / "verify_results.csv", index=False)
print(out.to_string())
print("\nDISCREPANCIES:", int((out.verdict == "DISCREPANCY").sum()), "of", len(out))
